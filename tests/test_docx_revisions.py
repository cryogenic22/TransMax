"""
TMX-3700 — DOCX revision-mark detection.

Confirms `DocxIngestionService` v2 attaches `meta.revisions` to blocks that
contain tracked-change marks (<w:ins>, <w:del>, <w:moveFrom>, <w:moveTo>),
with author + date provenance. Body text remains FINAL — deleted content
is excluded, inserted content is part of the segment.

Tests fabricate DOCX files in temp paths by injecting raw revision XML into
python-docx output, then run the ingester and assert metadata.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from lxml import etree

from app.services.docx_ingestion import DocxIngestionService, _collect_revisions
from app.services.docx_utils import (
    AUTHOR_ATTR, DATE_ATTR, DEL_NS, INS_NS, MOVE_FROM_NS, MOVE_TO_NS,
    PNS, TNS, W,
)


# ── Helpers — construct DOCX with tracked changes ──────────────────────


W_NS = {"w": W}


def _make_run(text: str) -> etree._Element:
    """Build a <w:r><w:t>text</w:t></w:r>."""
    r = etree.SubElement(etree.Element("dummy"), f"{{{W}}}r")
    t = etree.SubElement(r, TNS)
    t.text = text
    # Detach r from dummy; lxml needs us to remove it from a parent.
    r.getparent().remove(r)
    return r


def _make_ins(text: str, author: str = "Reviewer A", date: str = "2026-04-01T10:00:00Z") -> etree._Element:
    """Build <w:ins w:author=... w:date=...><w:r><w:t>text</w:t></w:r></w:ins>."""
    ins = etree.SubElement(etree.Element("dummy"), INS_NS)
    ins.set(AUTHOR_ATTR, author)
    ins.set(DATE_ATTR, date)
    ins.append(_make_run(text))
    ins.getparent().remove(ins)
    return ins


def _make_del(text: str, author: str = "Reviewer B", date: str = "2026-04-02T11:00:00Z") -> etree._Element:
    """Build <w:del w:author=... w:date=...><w:r><w:delText>text</w:delText></w:r></w:del>."""
    delete = etree.SubElement(etree.Element("dummy"), DEL_NS)
    delete.set(AUTHOR_ATTR, author)
    delete.set(DATE_ATTR, date)
    r = etree.SubElement(delete, f"{{{W}}}r")
    dtext = etree.SubElement(r, f"{{{W}}}delText")
    dtext.text = text
    delete.getparent().remove(delete)
    return delete


def _new_doc(tmp_path: Path) -> Path:
    """Create a minimal DOCX in tmp_path. Returns the path."""
    doc = DocxDocument()
    doc.add_paragraph("Plain paragraph.")
    p = tmp_path / "doc.docx"
    doc.save(str(p))
    return p


# ── Unit tests for _collect_revisions ─────────────────────────────────


def test_collect_revisions_returns_none_for_clean_element() -> None:
    p = etree.Element(PNS)
    r = etree.SubElement(p, f"{{{W}}}r")
    t = etree.SubElement(r, TNS)
    t.text = "no revisions here"
    assert _collect_revisions(p) is None


def test_collect_revisions_detects_insertion() -> None:
    p = etree.Element(PNS)
    p.append(_make_ins("inserted content", author="Dr. Reviewer", date="2026-04-01T10:00:00Z"))
    rev = _collect_revisions(p)
    assert rev is not None
    assert rev["has_insertions"] is True
    assert rev["has_deletions"] is False
    assert rev["authors"] == ["Dr. Reviewer"]
    assert rev["dates"] == ["2026-04-01T10:00:00Z"]


def test_collect_revisions_detects_deletion() -> None:
    p = etree.Element(PNS)
    p.append(_make_del("deleted content", author="Auditor", date="2026-04-02T11:00:00Z"))
    rev = _collect_revisions(p)
    assert rev is not None
    assert rev["has_deletions"] is True
    assert rev["has_insertions"] is False
    assert rev["authors"] == ["Auditor"]


def test_collect_revisions_dedupes_authors_and_dates() -> None:
    p = etree.Element(PNS)
    p.append(_make_ins("one", author="A", date="D1"))
    p.append(_make_del("two", author="A", date="D1"))  # same author + date
    p.append(_make_ins("three", author="B", date="D2"))
    rev = _collect_revisions(p)
    assert rev is not None
    assert rev["authors"] == ["A", "B"]
    assert rev["dates"] == ["D1", "D2"]


def test_collect_revisions_detects_moves() -> None:
    p = etree.Element(PNS)
    move = etree.SubElement(p, MOVE_TO_NS)
    move.set(AUTHOR_ATTR, "Mover")
    move.set(DATE_ATTR, "2026-04-03T12:00:00Z")
    r = etree.SubElement(move, f"{{{W}}}r")
    t = etree.SubElement(r, TNS)
    t.text = "moved here"
    rev = _collect_revisions(p)
    assert rev is not None
    assert rev["has_moves"] is True


# ── End-to-end through DocxIngestionService ───────────────────────────


def test_ingestion_clean_doc_has_no_revisions_meta(tmp_path: Path) -> None:
    path = _new_doc(tmp_path)
    blocks = DocxIngestionService().extract_blocks(str(path))
    # At least one paragraph block, no revisions metadata anywhere.
    assert any(b["type"] == "Paragraph" for b in blocks)
    for block in blocks:
        assert "revisions" not in block.get("meta", {}), (
            f"clean doc unexpectedly had revisions meta on block {block!r}"
        )


def test_ingestion_doc_with_insertion_carries_metadata(tmp_path: Path) -> None:
    """Build a doc with one inserted run, ingest, assert metadata captured."""
    doc = DocxDocument()
    para = doc.add_paragraph("Take the medication ")
    # Inject <w:ins> at the end of the paragraph's XML.
    p_elem = para._element
    p_elem.append(_make_ins("twice daily", author="Dr. Reviewer", date="2026-04-01T10:00:00Z"))
    path = tmp_path / "ins.docx"
    doc.save(str(path))

    blocks = DocxIngestionService().extract_blocks(str(path))
    para_blocks = [b for b in blocks if b["type"] == "Paragraph"]
    assert para_blocks, "expected at least one paragraph block"
    target = para_blocks[0]
    assert "revisions" in target["meta"]
    rev = target["meta"]["revisions"]
    assert rev["has_insertions"] is True
    assert rev["has_deletions"] is False
    assert rev["authors"] == ["Dr. Reviewer"]
    # Body text is FINAL: insertion is included.
    assert "twice daily" in target["text"]


def test_ingestion_doc_with_deletion_excludes_deleted_text(tmp_path: Path) -> None:
    """A deleted run's text must NOT appear in the segment, but metadata is captured."""
    doc = DocxDocument()
    para = doc.add_paragraph("Take medication every ")
    p_elem = para._element
    p_elem.append(_make_del("six hours", author="Auditor", date="2026-04-02T11:00:00Z"))
    # Append a final inserted run so we know the paragraph survives.
    p_elem.append(_make_ins("eight hours", author="Auditor", date="2026-04-02T11:00:00Z"))
    path = tmp_path / "del.docx"
    doc.save(str(path))

    blocks = DocxIngestionService().extract_blocks(str(path))
    para_blocks = [b for b in blocks if b["type"] == "Paragraph"]
    assert para_blocks
    target = para_blocks[0]
    assert "revisions" in target["meta"]
    rev = target["meta"]["revisions"]
    assert rev["has_deletions"] is True
    assert rev["has_insertions"] is True
    # Deleted text "six hours" should NOT be in the segment.
    assert "six hours" not in target["text"], (
        f"Deleted text appeared in segment: {target['text']!r}"
    )
    # Inserted text "eight hours" SHOULD be in the segment.
    assert "eight hours" in target["text"]


def test_ingestion_multiple_authors_merged(tmp_path: Path) -> None:
    doc = DocxDocument()
    para = doc.add_paragraph("Final text reads as: ")
    p_elem = para._element
    p_elem.append(_make_ins("one", author="A", date="2026-04-01T10:00:00Z"))
    p_elem.append(_make_ins(" and two", author="B", date="2026-04-02T11:00:00Z"))
    path = tmp_path / "multi.docx"
    doc.save(str(path))

    blocks = DocxIngestionService().extract_blocks(str(path))
    para_blocks = [b for b in blocks if b["type"] == "Paragraph"]
    assert para_blocks
    target = para_blocks[0]
    rev = target["meta"]["revisions"]
    assert set(rev["authors"]) == {"A", "B"}
    assert len(rev["authors"]) == 2  # deduped, ordered by occurrence
    assert rev["authors"] == ["A", "B"]
