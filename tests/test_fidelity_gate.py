"""TMX-FIDELITY-GATE — export refuses to ship a structurally degraded DOCX."""
from __future__ import annotations

from docx import Document

from app.services.document_export import DocumentExportService
from app.services.fidelity import structural_loss


def test_structural_loss_flags_dropped_skeleton():
    src = {"tables": 2, "images": 1, "headings": 3, "bold_runs": 5, "italic_runs": 2}
    out = {"tables": 0, "images": 1, "headings": 3, "bold_runs": 0, "italic_runs": 0}
    lost = structural_loss(src, out)
    assert any("tables: 2 → 0" in x for x in lost)
    # bold/italic shifts are NOT gate-failing (text redistribution is legitimate).
    assert not any("bold" in x or "italic" in x for x in lost)


def test_structural_loss_none_when_preserved():
    fp = {"tables": 2, "images": 1, "headings": 3, "bold_runs": 5, "italic_runs": 2}
    assert structural_loss(fp, dict(fp)) == []


def _rich_docx(path):
    d = Document()
    d.add_heading("A Structured Title", level=0)
    d.add_heading("Section One", level=1)
    d.add_paragraph("Some body text that is long enough to be a real paragraph.")
    t = d.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            t.cell(r, c).text = f"cell {r}{c}"
    d.save(str(path))
    return d


def _identity_segments(doc):
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if cell.text.strip():
                    texts.append(cell.text.strip())
    return [{"order_index": i, "source_text": t, "translated_text": t} for i, t in enumerate(texts)]


def test_faithful_roundtrip_passes_the_gate(tmp_path):
    """A correct round-trip preserves the skeleton ⇒ no FidelityError (no false positive)."""
    p = tmp_path / "m.docx"
    d = _rich_docx(p)
    buf = DocumentExportService().export_docx(str(p), _identity_segments(d))
    out = Document(buf)
    # Skeleton survived.
    assert len(out.tables) == 1
    headings = sum(1 for pp in out.paragraphs
                   if pp.style and (pp.style.name.startswith("Heading") or pp.style.name in ("Title", "Subtitle")))
    assert headings == 2
