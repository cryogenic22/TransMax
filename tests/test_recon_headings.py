"""TMX-3712-RECON-headings — PDF export applies heading styles from element_type.

The NEJM manuscript came back as 299 flat 'Normal' paragraphs because the
PDF flat-export path ignored the element_type the ingester recorded. Now Title/
Header blocks become real DOCX Title/Heading styles.
"""
from __future__ import annotations

from docx import Document

from app.services.document_export import DocumentExportService
from app.services.fidelity import structure_fingerprint


def _segs():
    # element_type mirrors what the ingester stores (unstructured categories).
    return [
        {"order_index": 0, "source_text": "BNT162b2 Vaccine", "translated_text": "BNT162b2 Vaccine", "element_type": "Title"},
        {"order_index": 1, "source_text": "BACKGROUND", "translated_text": "BACKGROUND", "element_type": "Header"},
        {"order_index": 2, "source_text": "Safe vaccines are needed.", "translated_text": "Safe vaccines are needed.", "element_type": "NarrativeText"},
        {"order_index": 3, "source_text": "RESULTS", "translated_text": "RESULTS", "element_type": "SectionHeader"},
        {"order_index": 4, "source_text": "Efficacy was 95%.", "translated_text": "Efficacy was 95%.", "element_type": "Sentence"},
    ]


def test_pdf_export_applies_heading_styles(tmp_path):
    buf = DocumentExportService().export_docx("", _segs(), force_new=True)
    doc = Document(buf)

    styles = [(p.style.name if p.style else None, p.text) for p in doc.paragraphs if p.text.strip()]
    by_text = {t: s for s, t in styles}

    assert by_text["BNT162b2 Vaccine"] == "Title"
    assert by_text["BACKGROUND"].startswith("Heading")
    assert by_text["RESULTS"].startswith("Heading")
    assert by_text["Safe vaccines are needed."] == "Normal"

    # And the fidelity fingerprint now sees real headings (was 0 before).
    fp = structure_fingerprint(doc)
    assert fp["headings"] == 3  # Title + 2 section headers


def test_plain_blocks_still_normal(tmp_path):
    segs = [
        {"order_index": 0, "source_text": "Just body text here.", "translated_text": "Just body text here.", "element_type": "NarrativeText"},
        {"order_index": 1, "source_text": "More body.", "translated_text": "More body.", "element_type": None},
    ]
    doc = Document(DocumentExportService().export_docx("", segs, force_new=True))
    assert structure_fingerprint(doc)["headings"] == 0
