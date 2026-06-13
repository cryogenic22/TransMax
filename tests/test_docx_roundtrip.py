"""
Round-trip tests for DOCX ingestion and export.
Verifies that text extracted by DocxIngestionService is correctly replaced by DocumentExportService.
"""
import os
import pytest
from docx import Document as DocxDocument
from docx.table import _Cell

from app.services.docx_ingestion import DocxIngestionService
from app.services.document_export import DocumentExportService


@pytest.fixture
def ingestion():
    return DocxIngestionService()


@pytest.fixture
def exporter():
    return DocumentExportService()


# --- Helper to save and reload ---

def save_docx(doc, tmp_path, name="test.docx"):
    path = os.path.join(str(tmp_path), name)
    doc.save(path)
    return path


def build_segments(blocks, translate_fn=None):
    """Build segment dicts from ingestion blocks, optionally applying a translation function."""
    segments = []
    for i, block in enumerate(blocks):
        translated = translate_fn(block["text"]) if translate_fn else f"TR:{block['text']}"
        segments.append({
            "order_index": i + 1,
            "source_text": block["text"],
            "translated_text": translated,
            "element_type": block["type"],
            "element_meta": block.get("meta"),
        })
    return segments


# --- Tests ---


class TestBodyParagraphs:
    def test_paragraphs_roundtrip(self, tmp_path, ingestion, exporter):
        """Body paragraphs are extracted and replaced in order."""
        doc = DocxDocument()
        doc.add_paragraph("Hello world")
        doc.add_paragraph("Second paragraph")
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        assert len(blocks) == 2
        assert blocks[0]["text"] == "Hello world"
        assert blocks[0]["type"] == "Paragraph"
        assert blocks[1]["text"] == "Second paragraph"

        segments = build_segments(blocks)
        buf = exporter.export_docx(path, segments)

        # Reload and verify
        result = DocxDocument(buf)
        texts = [p.text for p in result.paragraphs if p.text.strip()]
        assert texts == ["TR:Hello world", "TR:Second paragraph"]

    def test_empty_paragraphs_skipped(self, tmp_path, ingestion):
        """Empty paragraphs are not included in blocks."""
        doc = DocxDocument()
        doc.add_paragraph("Content")
        doc.add_paragraph("")  # empty
        doc.add_paragraph("More content")
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        assert len(blocks) == 2
        assert [b["text"] for b in blocks] == ["Content", "More content"]


class TestTablesInterleaved:
    def test_paragraph_table_order(self, tmp_path, ingestion, exporter):
        """Paragraphs and tables are extracted in document order (interleaved)."""
        doc = DocxDocument()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Cell A"
        doc.add_paragraph("After table")
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        types = [b["type"] for b in blocks]
        texts = [b["text"] for b in blocks]

        assert types == ["Paragraph", "TableCell", "Paragraph"]
        assert texts == ["Before table", "Cell A", "After table"]

        # Export round-trip
        segments = build_segments(blocks)
        buf = exporter.export_docx(path, segments)
        result = DocxDocument(buf)

        # Check paragraph texts
        para_texts = [p.text for p in result.paragraphs if p.text.strip()]
        assert "TR:Before table" in para_texts
        assert "TR:After table" in para_texts

    def test_multi_cell_table(self, tmp_path, ingestion):
        """Multiple cells in a table are each extracted."""
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "R0C0"
        table.cell(0, 1).text = "R0C1"
        table.cell(1, 0).text = "R1C0"
        table.cell(1, 1).text = "R1C1"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        cell_blocks = [b for b in blocks if b["type"] == "TableCell"]
        assert len(cell_blocks) == 4
        assert [b["text"] for b in cell_blocks] == ["R0C0", "R0C1", "R1C0", "R1C1"]

    def test_table_cell_text_is_translated_on_export(self, tmp_path, ingestion, exporter):
        """Translated text must land INSIDE single-paragraph table cells.

        Regression: ``_replace_cell_text`` cleared the paragraph it had just
        written because it compared python-docx proxy identity
        (``p is not paras[0]``) instead of element identity (``p._p is ...``).
        A fresh proxy is built per ``cell.paragraphs`` access, so every
        single-paragraph cell — the common case in pharma SmPC/CSR tables —
        came back EMPTY while the table grid survived (FidelityGate passed).
        """
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        for (r, c) in [(0, 0), (0, 1), (1, 0), (1, 1)]:
            table.cell(r, c).text = f"R{r}C{c}"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        buf = exporter.export_docx(path, build_segments(blocks))
        result = DocxDocument(buf)

        cells = result.tables[0].rows
        got = [cell.text for row in cells for cell in row.cells]
        assert got == ["TR:R0C0", "TR:R0C1", "TR:R1C0", "TR:R1C1"], (
            f"table cells not translated correctly: {got!r}"
        )
        assert "" not in got, "single-paragraph cell came back empty (the regression)"


class TestMergedCells:
    def test_merged_cell_dedup(self, tmp_path, ingestion):
        """Merged cells are only extracted once."""
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Merged"
        # Merge cells (0,0) and (0,1)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(1, 0).text = "Bottom left"
        table.cell(1, 1).text = "Bottom right"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        cell_blocks = [b for b in blocks if b["type"] == "TableCell"]
        texts = [b["text"] for b in cell_blocks]
        # "Merged" should appear only once
        assert texts.count("Merged") == 1
        assert "Bottom left" in texts
        assert "Bottom right" in texts


class TestNestedTables:
    def test_nested_table_extraction(self, tmp_path, ingestion, exporter):
        """Nested tables are extracted with nested=True meta."""
        doc = DocxDocument()
        outer = doc.add_table(rows=1, cols=1)
        outer_cell = outer.cell(0, 0)
        outer_cell.text = "Outer cell"
        # Add nested table inside the cell
        inner = outer_cell.add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "Inner cell"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        cell_blocks = [b for b in blocks if b["type"] == "TableCell"]
        assert len(cell_blocks) == 2
        # First should be outer (nested=False), second inner (nested=True)
        assert cell_blocks[0]["text"] == "Outer cell"
        assert cell_blocks[0]["meta"]["nested"] is False
        assert cell_blocks[1]["text"] == "Inner cell"
        assert cell_blocks[1]["meta"]["nested"] is True


class TestHeaders:
    def test_header_extraction(self, tmp_path, ingestion, exporter):
        """Headers are extracted and can be round-tripped."""
        doc = DocxDocument()
        doc.add_paragraph("Body text")
        section = doc.sections[0]
        header = section.header
        header.is_linked_to_previous = False
        header.paragraphs[0].text = "My Header"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        header_blocks = [b for b in blocks if b["type"] == "Header"]
        assert len(header_blocks) == 1
        assert header_blocks[0]["text"] == "My Header"
        assert header_blocks[0]["meta"]["variant"] == "default"

        # Export round-trip
        segments = build_segments(blocks)
        buf = exporter.export_docx(path, segments)
        result = DocxDocument(buf)
        hdr_text = result.sections[0].header.paragraphs[0].text
        assert hdr_text == "TR:My Header"


class TestFooters:
    def test_footer_extraction(self, tmp_path, ingestion, exporter):
        """Footers are extracted and can be round-tripped."""
        doc = DocxDocument()
        doc.add_paragraph("Body text")
        section = doc.sections[0]
        footer = section.footer
        footer.is_linked_to_previous = False
        footer.paragraphs[0].text = "My Footer"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        footer_blocks = [b for b in blocks if b["type"] == "Footer"]
        assert len(footer_blocks) == 1
        assert footer_blocks[0]["text"] == "My Footer"

        # Export round-trip
        segments = build_segments(blocks)
        buf = exporter.export_docx(path, segments)
        result = DocxDocument(buf)
        ftr_text = result.sections[0].footer.paragraphs[0].text
        assert ftr_text == "TR:My Footer"


class TestOrderConsistency:
    def test_ingestion_export_order_matches(self, tmp_path, ingestion, exporter):
        """The order of blocks from ingestion matches the export traversal order."""
        doc = DocxDocument()
        doc.add_paragraph("Para 1")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Table cell"
        doc.add_paragraph("Para 2")

        section = doc.sections[0]
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = "Header text"
        section.footer.is_linked_to_previous = False
        section.footer.paragraphs[0].text = "Footer text"
        path = save_docx(doc, tmp_path)

        blocks = ingestion.extract_blocks(path)
        # Body first, then headers, then footers
        types = [b["type"] for b in blocks]
        body_end = max(
            i for i, t in enumerate(types) if t in ("Paragraph", "TableCell")
        )
        header_indices = [i for i, t in enumerate(types) if t == "Header"]
        footer_indices = [i for i, t in enumerate(types) if t == "Footer"]

        # Headers come after body
        if header_indices:
            assert min(header_indices) > body_end
        # Footers come after headers
        if footer_indices and header_indices:
            assert min(footer_indices) > max(header_indices)

        # Round-trip: verify all translations applied
        segments = build_segments(blocks)
        buf = exporter.export_docx(path, segments)
        result = DocxDocument(buf)

        para_texts = [p.text for p in result.paragraphs if p.text.strip()]
        assert "TR:Para 1" in para_texts
        assert "TR:Para 2" in para_texts
