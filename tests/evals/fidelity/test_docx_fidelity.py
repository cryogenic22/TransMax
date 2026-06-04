"""FIDELITY-EVAL — document structural-fidelity golden set + eval.

Two guarantees, both independent of translation language:

  1. **Identity round-trip**: export the SAME text back into a structured DOCX
     and assert tables / figures / headings / formatting survive. This is the
     cheap, powerful check — any structural drift is a fidelity defect
     regardless of language.
  2. **Regression detector**: the comparator must FLAG the NEJM failure mode
     (13 tables → 0, figures dropped, headings flattened) as a loss with a low
     score — proving the eval would have caught the bad manuscript output.

Golden corpus: a synthetic rich manuscript fixture (title + headings + a data
table + bold/italic runs). Extend with real-format docs over time.
"""
from __future__ import annotations

from docx import Document

from app.services.document_export import DocumentExportService
from app.services.fidelity import (
    compare_structure,
    structure_fingerprint,
)


def _build_rich_manuscript(path) -> Document:
    """A structured scientific-manuscript fixture (mirrors the NEJM shape)."""
    d = Document()
    d.add_heading("Safety and Efficacy of the BNT162b2 mRNA Covid-19 Vaccine", level=0)  # Title
    d.add_heading("BACKGROUND", level=1)
    p = d.add_paragraph("Severe acute respiratory syndrome ")
    p.add_run("coronavirus 2").italic = True
    p.add_run(" emerged in 2019.")
    p2 = d.add_paragraph()
    p2.add_run("Safe and effective vaccines are urgently needed.").bold = True
    d.add_heading("RESULTS", level=1)
    t = d.add_table(rows=2, cols=3)
    cells = [("Group", "N", "Efficacy"), ("BNT162b2", "21720", "95%")]
    for r_i, row in enumerate(cells):
        for c_i, val in enumerate(row):
            t.cell(r_i, c_i).text = val
    d.save(str(path))
    return d


def _identity_segments(doc) -> list[dict]:
    """Build identity segments (translated == source) for every text element."""
    texts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            texts.append(p.text.strip())
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if cell.text.strip():
                    texts.append(cell.text.strip())
    return [
        {"order_index": i, "source_text": t, "translated_text": t}
        for i, t in enumerate(texts)
    ]


def test_identity_roundtrip_preserves_structure(tmp_path):
    """Export the same text back → tables, figures, headings, formatting intact."""
    src_path = tmp_path / "manuscript.docx"
    src_doc = _build_rich_manuscript(src_path)
    src_fp = structure_fingerprint(src_doc)

    segments = _identity_segments(src_doc)
    buf = DocumentExportService().export_docx(str(src_path), segments)
    out_doc = Document(buf)
    out_fp = structure_fingerprint(out_doc)

    report = compare_structure(src_fp, out_fp)
    # The critical structure must be fully preserved — no dropped tables/figures/headings.
    assert out_fp["tables"] == src_fp["tables"], report.losses
    assert out_fp["images"] == src_fp["images"], report.losses
    assert out_fp["headings"] == src_fp["headings"], report.losses
    assert report.score >= 90.0, f"score={report.score} losses={report.losses}"


def test_comparator_flags_the_nejm_flattening():
    """The exact NEJM regression: a structured source vs a flattened output is
    caught as a loss with a low score (the eval would have stopped it)."""
    source = {"paragraphs": 299, "tables": 13, "images": 5, "headings": 12,
              "bold_runs": 40, "italic_runs": 20}
    flattened = {"paragraphs": 299, "tables": 0, "images": 0, "headings": 0,
                 "bold_runs": 0, "italic_runs": 0}
    report = compare_structure(source, flattened)
    assert report.ok is False
    assert report.score < 10.0, report.score
    joined = " ".join(report.losses)
    assert "tables: 13 → 0" in joined
    assert "images: 5 → 0" in joined
    assert "headings: 12 → 0" in joined


def test_clean_when_structure_preserved():
    fp = {"paragraphs": 10, "tables": 2, "images": 1, "headings": 3,
          "bold_runs": 4, "italic_runs": 2}
    report = compare_structure(fp, dict(fp))
    assert report.ok is True
    assert report.score == 100.0
    assert report.losses == []


def test_added_paragraphs_not_penalized():
    """A more-verbose target language may add paragraphs — not a loss."""
    src = {"paragraphs": 10, "tables": 1, "images": 0, "headings": 2,
           "bold_runs": 0, "italic_runs": 0}
    out = {"paragraphs": 14, "tables": 1, "images": 0, "headings": 2,
           "bold_runs": 0, "italic_runs": 0}
    report = compare_structure(src, out)
    assert report.ok is True
    assert report.score == 100.0
